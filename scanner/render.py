"""Render a tailored resume / cover letter to DOCX and PDF. One ATS-friendly single-column layout."""
import io
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


def _lines(doc: dict):
    """Resume as a flat list of (kind, text) — shared by both renderers."""
    yield "name", doc["name"]
    yield "contact", " · ".join(doc["contact"])
    yield "heading", "Summary"
    yield "para", doc["summary"]
    yield "heading", "Experience"
    for e in doc["experience"]:
        place = " · ".join(x for x in (e["company"], e.get("location")) if x)
        yield "role", f"{e['title']} — {place} · {e['dates']}"
        for b in e["bullets"]:
            yield "bullet", b["text"]
    if doc["projects"]:
        yield "heading", "Projects"
        for p in doc["projects"]:
            yield "role", p["name"] + (f" — {p['url']}" if p["url"] else "")
            for b in p["bullets"]:
                yield "bullet", b
    yield "heading", "Education"
    for e in doc["education"]:
        yield "para", e
    yield "heading", "Skills"
    yield "para", ", ".join(doc["skills"])
    if doc["interests"]:
        yield "para", f"Interests: {doc['interests']}"


def resume_docx(doc: dict) -> bytes:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    for kind, text in _lines(doc):
        if kind == "name":
            d.add_heading(text, level=0)
        elif kind == "heading":
            d.add_heading(text, level=1)
        elif kind == "role":
            d.add_paragraph().add_run(text).bold = True
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(flowables) -> bytes:
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=15 * mm, bottomMargin=15 * mm).build(flowables)
    return buf.getvalue()


def _styles():
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5)
    return {
        "name": ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=16, leading=20),
        "contact": ParagraphStyle("contact", parent=body, fontSize=8.5, spaceAfter=4),
        "heading": ParagraphStyle("heading", parent=body, fontName="Helvetica-Bold", fontSize=11,
                                  spaceBefore=8, spaceAfter=3),
        "role": ParagraphStyle("role", parent=body, fontName="Helvetica-Bold", spaceBefore=4),
        "para": body,
        "bullet": body,
    }


def resume_pdf(doc: dict) -> bytes:
    st = _styles()
    flow, bullets = [], []

    def flush():
        if bullets:
            flow.append(ListFlowable([ListItem(b, leftIndent=10) for b in bullets], bulletType="bullet",
                                     start="•", leftIndent=10, bulletFontSize=8))
            bullets.clear()

    for kind, text in _lines(doc):
        p = Paragraph(escape(text), st[kind])
        if kind == "bullet":
            bullets.append(p)
        else:
            flush()
            flow.append(p)
    flush()
    return _pdf(flow)


def letter_docx(text: str) -> bytes:
    d = Document()
    d.styles["Normal"].font.size = Pt(11)
    for para in text.split("\n\n"):
        d.add_paragraph(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def letter_pdf(text: str) -> bytes:
    body = ParagraphStyle("letter", parent=getSampleStyleSheet()["Normal"], fontName="Helvetica",
                          fontSize=11, leading=15, spaceAfter=10)
    return _pdf([Paragraph(escape(p.strip()).replace("\n", "<br/>"), body) for p in text.split("\n\n")]
                + [Spacer(1, 1)])
